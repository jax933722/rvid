"""Unit tests for the CSV / JSON / XLSX exporters."""

from __future__ import annotations

import csv
import io
import json

from openpyxl import load_workbook

from bise.application.dto.export_dto import ExportFormat
from bise.infrastructure.export.exporters import build_exporter

COLUMNS = ("company_id", "display_name", "technologies", "employee_count")
ROWS = [
    {
        "company_id": 1,
        "display_name": "Acme",
        "technologies": "WordPress; Shopify",
        "employee_count": 12,
    },
    {"company_id": 2, "display_name": "Beta", "technologies": "", "employee_count": None},
]


def test_csv_exporter_roundtrips() -> None:
    exporter = build_exporter(ExportFormat.CSV)
    assert exporter.extension == "csv"
    data = exporter.serialize(ROWS, COLUMNS).decode("utf-8-sig")
    reader = list(csv.DictReader(io.StringIO(data)))
    assert [r["display_name"] for r in reader] == ["Acme", "Beta"]
    assert reader[0]["technologies"] == "WordPress; Shopify"
    assert reader[1]["employee_count"] == ""  # None -> blank


def test_json_exporter_roundtrips() -> None:
    exporter = build_exporter(ExportFormat.JSON)
    payload = json.loads(exporter.serialize(ROWS, COLUMNS))
    assert payload[0]["company_id"] == 1
    assert payload[1]["employee_count"] is None
    assert set(payload[0].keys()) == set(COLUMNS)


def test_xlsx_exporter_produces_openable_workbook() -> None:
    exporter = build_exporter(ExportFormat.XLSX)
    assert exporter.extension == "xlsx"
    workbook = load_workbook(io.BytesIO(exporter.serialize(ROWS, COLUMNS)))
    sheet = workbook.active
    header = [c.value for c in sheet[1]]
    assert header == list(COLUMNS)
    assert sheet.cell(row=2, column=2).value == "Acme"
